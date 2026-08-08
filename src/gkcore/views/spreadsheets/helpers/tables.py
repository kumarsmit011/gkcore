# Spreadsheet libraries
from openpyxl.styles import Font, Alignment


def add_table(sheet, org_title, page_title, fields, values, table_first_row):
    """ Generic helper to add tables to spreadsheets. This helper returns, the updated
    sheet and the last row number of the table.

    Parameters:
    - sheet: Instance of Openpyxl Worksheet
    - org_title: Organisation title
    - page_title: Page title
    - fields: Field list with key to map to value. Fields will be of format,
      [{"key": "", "label": "", "width": ""}, ...]
    - value: Values for table rows
    - table_first_row: Row where table starts
    """

    sheet.merge_cells("A1:H2")
    sheet["A1"].font = Font(size="16", bold=True)
    sheet["A1"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["A1"] = org_title

    sheet.merge_cells("A3:H3")
    # Title of the sheet and width of columns are set.
    sheet.title = page_title
    sheet["A3"] = page_title
    sheet["A3"].alignment = Alignment(horizontal="center", vertical="center")
    sheet["A3"].font = Font(size="13", bold=True)

    table_last_row = 0
    for position, field in enumerate(fields):
        if not field.get("key"):
            continue
        column = convert_number_to_column(position+1)
        sheet.column_dimensions[column].width = field.get("width") or 20
        sheet[f"{column}{table_first_row}"] = field.get("label") or field["key"].title()
        sheet[f"{column}{table_first_row}"].font = Font(size="13", bold=True)
        row = table_first_row
        for item in values:
            row += 1
            if item.get(field["key"]):
                sheet[f"{column}{row}"] = item[field["key"]]
        table_last_row = row
    return sheet, table_last_row


def convert_number_to_column(number, character=""):
    if number <= 26:
        character += chr(ord('@')+number)
        return character
    character += "Z"
    number -= 26
    convert_number_to_column(number, character)
