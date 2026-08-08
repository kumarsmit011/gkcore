import json

from gkcore.views.spreadsheets.helpers.tables import add_table, convert_number_to_column
from pyramid.response import Response
from pyramid.request import Request

import openpyxl
from openpyxl.styles import DEFAULT_FONT, Font

# from io import BytesIO
import io

def view_register_xlsx_generator(request):
    """ Generates response with spreadsheet file generated from view register_xlsx API.

    """
    calculatefrom = request.params["from"]
    calculateto = request.params["to"]
    title = request.params["title"]
    fields = json.loads(request.params["fields"])
    header = {"gktoken": request.headers["gktoken"]}
    fystart = str(request.params["fystart"])
    fyend = str(request.params["fyend"])
    orgname = str(request.params["orgname"])

    req = Request.blank(
        "/reports/registers?flag=0&calculatefrom=%s&calculateto=%s"
        % (calculatefrom, calculateto),
        headers=header,
    )

    response = json.loads(request.invoke_subrequest(req).text)["gkresult"]

    vouchers = response["vouchers"]

    org_title = f"{orgname.upper()} (FY: {fystart} to {fyend})"
    page_title = f"{title} from {fystart} to {fyend}"

    # A workbook is opened.
    DEFAULT_FONT.name = "Liberation Serif"
    wb = openpyxl.Workbook()
    sheet = wb.active
    table_first_row = 4

    tax_totals = response["tax_totals"]
    totals = {
        "custname": "TOTAL",
        **tax_totals,
        "amount": response["voucher_total"],
        "taxed": response["taxed_total"],
    }

    # Adding tax values to dictionary root
    vouchers = [
       {
           **voucher,
           **{tax['tax_str']:tax["tax_amount"] for tax in voucher["tax_data"]}
       } for voucher in vouchers
    ]

    vouchers.append(totals)

    sheet, last_row = add_table(
        sheet,
        org_title,
        page_title,
        fields,
        vouchers,
        table_first_row,
    )

    column = ""
    for column_no in range(1, len(fields)+1):
        column = convert_number_to_column(column_no)
        for row in range(1, last_row+1):
            sheet[f"{column}{row}"].number_format = '#,##0.00'

    for cell in sheet[last_row]:
        cell.font = Font(size="13", bold=True)
    sheet[f"{column}{last_row}"].font = Font(size="13", bold=True, u="doubleAccounting")

    output = io.BytesIO()
    wb.save(output)
    contents = output.getvalue()
    output.close()
    headerList = {
        "Content-Type": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        "Content-Length": len(contents),
        "Content-Disposition": "attachment; filename=report.xlsx",
        "X-Content-Type-Options": "nosniff",
        "Set-Cookie": "fileDownload=true; path=/ [;HttpOnly]",
    }
    return Response(contents, headerlist=list(headerList.items()))
